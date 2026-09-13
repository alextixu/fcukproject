# -*- coding: utf-8 -*-
"""週報簡報:從「猜漲跌」到「排順序」(5 頁版)。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(ROOT, "figs", "decile")
OUT = os.path.join(ROOT, "十分位多空_週報.pptx")
PAPER_URL = "https://onlinelibrary.wiley.com/doi/abs/10.1111/jofi.13268"

W, H = 13.333, 7.5
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
LGREY = RGBColor(0xF2, 0xF2, 0xF2)
ACC = RGBColor(0x00, 0x78, 0xD4)
NEG = RGBColor(0xD1, 0x34, 0x38)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft JhengHei"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)


def _f(run, size, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def blank():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = WHITE
    return s


def bar(s, top=0.42, h=0.62, color=INK, left=0.62, w=0.07):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                            Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def text(s, x, y, w, h, items, align=PP_ALIGN.LEFT, spacing=1.0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (t, sz, bd, col) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = t
        _f(r, sz, bd, col)
    return tb


def title(s, t, sub=None):
    bar(s, 0.44, 0.70)
    text(s, 0.95, 0.34, 11.9, 0.9, [(t, 30, True, INK)])
    if sub:
        text(s, 0.95, 1.16, 11.9, 0.5, [(sub, 14, False, GREY)])


def table(s, data, x, y, w, h, col_w=None, fs=13, head_fs=13, hi_rows=(),
          blue_rows=()):
    r, c = len(data), len(data[0])
    g = s.shapes.add_table(r, c, Inches(x), Inches(y), Inches(w),
                           Inches(h)).table
    if col_w:
        for i, cw in enumerate(col_w):
            g.columns[i].width = Inches(cw)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = g.cell(i, j)
            cell.text = str(val)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.09)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                col = WHITE if i == 0 else (
                    NEG if i in hi_rows else (
                        ACC if i in blue_rows else INK))
                bd = i == 0 or i in hi_rows or i in blue_rows
                if not p.runs:
                    p.add_run().text = ""
                for rr in p.runs:
                    _f(rr, head_fs if i == 0 else fs, bd, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (INK if i == 0 else
                                        (LGREY if i % 2 == 0 else WHITE))
    return g


def note(s, t, y=6.62, color=ACC, h=0.62, extra=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(y),
                            Inches(12.1), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = LGREY
    sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = t
    _f(r, 14, True, color)
    if extra:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.LEFT
        r2 = p2.add_run()
        r2.text = extra
        _f(r2, 13, False, GREY)


# ══════════════ 1 封面 ══════════════
s = blank()
text(s, 1.0, 2.45, 11.3, 1.4,
     [("從「猜漲跌」到「排順序」", 46, True, INK)], PP_ALIGN.CENTER)
text(s, 1.0, 3.65, 11.3, 0.8,
     [("Chart GCN 評估協定的轉型", 26, False, ACC)], PP_ALIGN.CENTER)
bar(s, 4.85, 0.06, ACC, 6.16, 1.0)
text(s, 0.9, 6.5, 6, 0.8, [("專題進度週報", 14, False, GREY),
                           ("2026-08-25", 14, False, GREY)])

# ══════════════ 2 相似點 ══════════════
s = blank()
title(s, "JKX 的方法與我們高度相似",
      "同一個大方向:不預設型態,讓卷積網路自己從價格圖表中學")
table(s, [
    ["面向", "我們的 Chart GCN", "JKX (JF 2023)"],
    ["核心命題", "圖表型態含有預測資訊", "同"],
    ["是否預設型態", "否(不用頭肩、黃金交叉)", "否(明確批評 ad hoc 型態)"],
    ["輸入原料", "日 OHLC + 成交量", "日 OHLC + 成交量"],
    ["表示法家族", "價格 → 二維張量", "價格 → 二維像素圖"],
    ["模型家族", "Conv2d × 3 + Attention", "Conv 區塊 × 2~4"],
    ["參數量", "38,130", "155K ~ 2.95M"],
    ["標籤形式", "未來漲跌二分類", "未來漲跌二分類"],
    ["損失函數", "交叉熵 + Adam", "交叉熵 + Adam"],
], 0.95, 2.05, 11.5, 4.2, [2.5, 4.6, 4.4], fs=13.5)
note(s, "方法論家族相同 —— 所以 JKX 的評估協定可以、也應該直接套用到我們的模型上。")

# ══════════════ 3 JKX 的分類結果 ══════════════
s = blank()
title(s, "JKX 論文的分類結果:準確率只有 53%",
      "Table 2,美股 CRSP 全樣本,1993–2000 訓練一次後凍結,2001–2019 樣本外")
table(s, [
    ["圖像長度", "預測 20 日", "預測 60 日"],
    ["5 日", "52.5%", "53.6%"],
    ["20 日", "53.3%", "53.2%"],
    ["60 日", "53.6%", "52.9%"],
    ["動量 MOM(對照)", "52.1%", "52.1%"],
], 0.95, 2.15, 5.4, 2.2, [2.2, 1.6, 1.6], fs=13.5)
table(s, [
    ["Chart GCN 論文報的(SZ-50)", "數值"],
    ["Accuracy", "69.26%"],
    ["Precision(漲 / 跌)", "65.76% / 72.26%"],
    ["F1(漲 / 跌)", "66.40% / 71.67%"],
    ["JKX 是否報 F1", "否 —— 財金領域不用"],
], 6.75, 2.15, 5.7, 2.2, [3.0, 2.7], fs=13.5, hi_rows=(1,))
text(s, 0.95, 4.7, 11.7, 1.7, [
    ("JKX 花了半頁論證「53% 已經夠用」:準確率每 +1%,擇時策略年化 Sharpe 約 +0.1。",
     18, False, INK),
    ("Chart GCN 宣稱的 69.26%,比 JKX 的天花板高出 16pp。", 21, True, NEG),
], spacing=1.35)
note(s, "所以他們的主證據不是準確率,而是下一頁的排序邏輯。")

# ══════════════ 4 排序邏輯 ══════════════
s = blank()
title(s, "他們的排序邏輯:十分位多空",
      "把「二分類 → 數對幾次」換成「機率 → 橫斷面排序」")
steps = [
    ("1", "取機率", "對當日每一檔股票輸出 softmax P(漲),而不是只取 argmax"),
    ("2", "排序", "把當日全部股票依 P(漲) 由低到高排隊"),
    ("3", "切十等分", "每一分位放同樣多檔股票"),
    ("4", "建立部位", "做多第 10 分位(最看好)、放空第 1 分位(最看壞),等權"),
    ("5", "持有並重排", "持有 h 個交易日後整批重新排序(不重疊)"),
]
y = 2.05
for n, t, d in steps:
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), Inches(y),
                            Inches(0.44), Inches(0.44))
    sh.fill.solid()
    sh.fill.fore_color.rgb = ACC
    sh.line.fill.background()
    p = sh.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = n
    _f(r, 15, True, WHITE)
    text(s, 1.60, y - 0.05, 2.0, 0.5, [(t, 18, True, INK)])
    text(s, 3.45, y - 0.03, 9.1, 0.5, [(d, 15, False, GREY)])
    y += 0.68
text(s, 0.95, 5.55, 11.7, 1.0, [
    ("為什麼比猜方向容易:大盤下跌時,第 10 分位 -2%、第 1 分位 -8% → "
     "準確率全錯,但 H-L = +6%。", 17, False, INK),
    ("多空對沖後大盤漲跌被消掉 —— Accuracy 量的是「會不會看大盤」,H-L 量的是「會不會選股」。",
     17, False, GREY),
], spacing=1.3)
note(s, "H-L 報酬 = 第 10 分位平均報酬 − 第 1 分位平均報酬。")

# ══════════════ 5 JKX 十分位結果 + 週轉率(最後一頁) ══════════════
s = blank()
title(s, "JKX 的十分位結果:報酬單調遞增",
      "Table 3,20 日圖 → 20 日報酬(I20/R20),等權,月頻再平衡")
table(s, [
    ["分位", "1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "H-L"],
    ["年化報酬 %", "-2", "5", "7", "9", "11", "11", "13", "14", "15", "18",
     "+21***"],
    ["Sharpe", "-0.12", "0.28", "0.39", "0.49", "0.59", "0.59", "0.74",
     "0.81", "0.87", "1.04", "2.16"],
], 0.62, 2.05, 12.1, 1.5, [1.9] + [0.85] * 10 + [1.35], fs=12, head_fs=12)
table(s, [
    ["同期對照(H-L,等權)", "年化報酬", "Sharpe", "月週轉率"],
    ["CNN I20/R20", "+21%", "2.16", "173%"],
    ["動量 MOM", "+7%", "0.25", "63%"],
    ["短期反轉 STR", "+11%", "0.55", "168%"],
    ["週反轉 WSTR", "+18%", "1.23", "167%"],
], 0.95, 3.85, 11.5, 2.0, [4.4, 2.4, 2.3, 2.4], fs=13, blue_rows=(1,))
note(s, "重點不是 Sharpe 2.16,而是前十欄「一路遞增、沒有亂跳」—— 這代表機率是連續有意義的。",
     y=5.95)

# 論文出處(最後一頁)
tb = s.shapes.add_textbox(Inches(0.62), Inches(6.72), Inches(12.1),
                          Inches(0.55))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
r0 = p.add_run()
r0.text = ("論文出處:Jiang, J., Kelly, B., & Xiu, D. (2023). "
           "(Re-)Imag(in)ing Price Trends. Journal of Finance, 78(6). ")
_f(r0, 12, False, GREY)
r1 = p.add_run()
r1.text = PAPER_URL
_f(r1, 12, False, ACC)
r1.hyperlink.address = PAPER_URL

prs.save(OUT)
print("已輸出:", OUT)
print("投影片數:", len(prs.slides._sldIdLst))
